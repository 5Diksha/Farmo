from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font


EXPORT_DIR = Path("exports")
EXPORT_DIR.mkdir(exist_ok=True)


def create_monthly_excel_report(
    month_name,
    year,
    work_records,
    payments,
    summary
):

    file_name = f"Farmo_Report_{month_name}_{year}.xlsx"

    file_path = EXPORT_DIR / file_name

    wb = Workbook()

    # =========================
    # SHEET 1 - WORK REGISTER
    # =========================

    ws1 = wb.active
    ws1.title = "Work Register"

    ws1.append([
        "Date",
        "Customer",
        "Work Type",
        "Quantity",
        "Advance Received"
    ])
    for cell in ws1[1]:
        cell.font = Font(bold=True)

    for row in work_records:

        if row["work_type"] in [
            "Rotavator",
            "Nangarat",
            "Fanadi"
        ]:
            quantity = f"{row['hours']}h {row['minutes']}m"
        else:
            quantity = f"{row['trips']} Trips"

        ws1.append([
            row["work_date"],
            row["customer_name"],
            row["work_type"],
            quantity,
            row["credit_given"]
        ])

    # =========================
    # SHEET 2 - PAYMENTS
    # =========================

    ws2 = wb.create_sheet("Payments")

    ws2.append([
        "Date",
        "Customer",
        "Amount",
        "Notes"
    ])
    for cell in ws2[1]:
        cell.font = Font(bold=True)

    for row in payments:

        ws2.append([
            row["payment_date"],
            row["customer_name"],
            row["amount"],
            row["notes"]
        ])

    # =========================
    # SHEET 3 - SUMMARY
    # =========================

    ws3 = wb.create_sheet("Monthly Summary")

    ws3.append(["Item", "Amount"])
    for cell in ws3[1]:
        cell.font = Font(bold=True)

    ws3.append([
        "Total Work Amount",
        summary["total_work_amount"]
    ])

    ws3.append([
        "Advance Received",
        summary["total_advance"]
    ])

    ws3.append([
        "Payments Received",
        summary["total_payments"]
    ])

    ws3.append([
        "Pending Balance",
        summary["pending_balance"]
    ])



    # =========================
    # SHEET 4 - WORK STATISTICS
    # =========================

    ws4 = wb.create_sheet("Work Statistics")

    ws4.append([
    "Work Type",
    "Quantity",
    "Earnings"
    ])

    for cell in ws4[1]:
        cell.font = Font(bold=True)

    rot_hours = 0
    rot_amount = 0

    nangarat_hours = 0
    nangarat_amount = 0

    fanadi_hours = 0
    fanadi_amount = 0

    trolley_trips = 0
    trolley_amount = 0

    tanker_trips = 0
    tanker_amount = 0

    for row in work_records:

      if row["work_type"] == "Rotavator":
        rot_hours += row["hours"] + (row["minutes"] / 60)
        rot_amount += row["amount"]

      elif row["work_type"] == "Nangarat":
        nangarat_hours += row["hours"] + (row["minutes"] / 60)
        nangarat_amount += row["amount"]

      elif row["work_type"] == "Fanadi":
        fanadi_hours += row["hours"] + (row["minutes"] / 60)
        fanadi_amount += row["amount"]

      elif row["work_type"] == "Trolley Trip":
        trolley_trips += row["trips"]
        trolley_amount += row["amount"]

      elif row["work_type"] == "Water Tanker Trip":
        tanker_trips += row["trips"]
        tanker_amount += row["amount"]

    ws4.append(["Rotavator", rot_hours, rot_amount])
    ws4.append(["Nangarat", nangarat_hours, nangarat_amount])
    ws4.append(["Fanadi", fanadi_hours, fanadi_amount])
    ws4.append(["Trolley Trip", trolley_trips, trolley_amount])
    ws4.append(["Water Tanker Trip", tanker_trips, tanker_amount])



    # DEBUG
    print("Workbook sheets:")
    print(wb.sheetnames)

    wb.save(file_path)

    print(f"Saved: {file_path}")

    return file_path

def create_customer_monthly_report(
    customer_name,
    month_name,
    year,
    work_records,
    payments,
    summary
):

    file_name = (
        f"{customer_name}_{month_name}_{year}.xlsx"
    )

    file_path = EXPORT_DIR / file_name

    wb = Workbook()

    # =========================
    # SHEET 1 - WORK SUMMARY
    # =========================

    ws1 = wb.active
    ws1.title = "Work Summary"

    ws1.append([
        "Work Type",
        "Quantity",
        "Amount"
    ])

    for cell in ws1[1]:
        cell.font = Font(bold=True)

    ws1.append([
        "Rotavator",
        summary["rot_hours"],
        summary["rot_amount"]
    ])

    ws1.append([
        "Nangarat",
        summary["nangarat_hours"],
        summary["nangarat_amount"]
    ])

    ws1.append([
        "Fanadi",
        summary["fanadi_hours"],
        summary["fanadi_amount"]
    ])

    ws1.append([
        "Trolley Trip",
        summary["trolley_trips"],
        summary["trolley_amount"]
    ])

    ws1.append([
        "Water Tanker Trip",
        summary["tanker_trips"],
        summary["tanker_amount"]
    ])

    # =========================
    # SHEET 2 - WORK HISTORY
    # =========================

    ws2 = wb.create_sheet("Work History")

    ws2.append([
        "Date",
        "Work Type",
        "Hours",
        "Minutes",
        "Trips",
        "Amount",
        "Advance Received"
    ])

    for cell in ws2[1]:
        cell.font = Font(bold=True)

    for row in work_records:

        ws2.append([
            row["work_date"],
            row["work_type"],
            row["hours"],
            row["minutes"],
            row["trips"],
            row["amount"],
            row["credit_given"]
        ])

    # =========================
    # SHEET 3 - PAYMENT HISTORY
    # =========================

    ws3 = wb.create_sheet("Payment History")

    ws3.append([
        "Date",
        "Amount",
        "Notes"
    ])

    for cell in ws3[1]:
        cell.font = Font(bold=True)

    for row in payments:

        ws3.append([
            row["payment_date"],
            row["amount"],
            row["notes"]
        ])

    # =========================
    # SHEET 4 - FINANCIAL SUMMARY
    # =========================

    ws4 = wb.create_sheet("Financial Summary")

    ws4.append(["Item", "Amount"])

    for cell in ws4[1]:
        cell.font = Font(bold=True)

    ws4.append([
        "Total Work Amount",
        summary["total_work_amount"]
    ])

    ws4.append([
        "Advance Received",
        summary["total_credit"]
    ])

    ws4.append([
        "Payments Received",
        summary["total_payments"]
    ])

    ws4.append([
        "Pending Balance",
        summary["balance"]
    ])

    wb.save(file_path)

    return file_path